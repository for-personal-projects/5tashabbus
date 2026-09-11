"""Ro'yxatdan o'tgan bolalarni Excel (xlsx) fayliga chiqarish."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER = ["№", "Familya", "Ism", "Sharif (otasining ismi)", "To'liq ism",
    "Tug'ilgan sana", "Jinsi", "Hujjat turi", "Hujjat seriyasi", "Hujjat raqami",
    "PINFL", "Yo'nalish", "Viloyat", "Tuman/shahar", "Mahalla (MFY)",
    "Sport yo'nalishi", "Sport turi", "Yosh kategoriya",
    "Qaysi raqamdan ro'yxatdan o'tgan", "Bolaning telefoni",
    "Ro'yxatdan o'tgan vaqti"]

# matn sifatida saqlanishi kerak bo'lgan ustun indekslari (0-asosli): 0/1 lar yeyilmasin
TEXTCOLS = {8, 9, 10, 18, 19}  # seriya, raqam, PINFL, telefonlar

FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def val(row: dict, key: str, default=""):
    d = row.get("data_json") or {}
    if key in d and d.get(key) not in (None, ""):
        return d.get(key, default)
    return row.get(key, default)


def build_excel(rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ro'yxat"

    ws.append(HEADER)
    for c in range(1, len(HEADER) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for i, r in enumerate(rows, 1):
        d = r.get("data_json") or {}
        gname = r.get("gendername") or d.get("gendername", "")
        docname = d.get("identity_name") or d.get("identitydocumentname", "")
        if not docname:
            import config
            docname = config.IDENTITY_DOCS.get(r.get("identity_id") or d.get("identity_id"), "")
        serp = str(r.get("series", "") or d.get("documentseries", ""))
        nump = str(r.get("number", "") or d.get("documentnumber", ""))
        pinf = str(d.get("pinfl", "") or "")
        sportcat = d.get("sportcat_name", "")
        sports = d.get("sport_names", "")
        agecat = d.get("agecat_name", "")
        region = d.get("region_name", "")
        mfy = d.get("mfy_name", "")
        oblast = d.get("oblast_name", "")
        phone = str(r.get("phone", "") or d.get("phonenumber", "") or d.get("mobilenumber", "") or "")
        mphone = str(d.get("mobilenumber", "") or d.get("phonenumber", "") or "")

        ws.append([
            i,
            d.get("familyname", ""), d.get("firstname", ""), d.get("lastname", ""),
            r.get("fullname", "") or d.get("fullname", "") or d.get("shortname", ""),
            d.get("dateofbirth", ""), gname,
            docname, serp, nump, pinf,
            d.get("initiativ_name", "") or r.get("initiativ", ""),
            oblast, region, mfy,
            sportcat, sports, agecat,
            phone, mphone,
            r.get("created_at", ""),
        ])

    # ustun kengliklari + formatlash
    for c in range(1, len(HEADER) + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = 14
    widths = {1: 6, 2: 18, 3: 16, 4: 20, 5: 26, 6: 13, 7: 10, 8: 22, 9: 13, 10: 13,
              11: 16, 12: 16, 13: 20, 14: 16, 15: 20, 16: 20, 17: 24, 18: 16,
              19: 22, 20: 18, 21: 20}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADER)):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = BORDER
            if cell.column - 1 in TEXTCOLS:
                cell.number_format = "@"   # matn: 0/1 bosh harflari yeyilmaydi
            if cell.column in (1, 6, 7, 9, 10, 11, 19, 20):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}1"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()