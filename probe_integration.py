"""Store(dan) -> Excel(export) yakuniy integratsion test."""
import asyncio
import sys
sys.path.insert(0, "d:/5tashabbus_bot")
import store
import exporter
from openpyxl import load_workbook


async def t():
    await store.init_db()
    uid = 555001
    athlete = {"fullname": "Toshtemirov Rahimjon", "familyname": "Toshtemirov",
               "firstname": "Rahimjon", "lastname": "Alisherovich",
               "dateofbirth": "10.10.2014", "gendername": "O'g'il",
               "documentseries": "I-HR", "documentnumber": "0456789",
               "pinfl": "11456987452036", "mobilenumber": ""}
    extra = {"initiativ_id": 2, "initiativ_name": "Maktab", "identity_id": 1,
             "identity_name": "Guvohnoma", "oblast_name": "Toshkent viloyati",
             "region_name": "Parkent tumani", "mfy_name": "77-NAVROZ",
             "sportcat_name": "Futbol", "sport_names": "Futbol 7x7",
             "agecat_id": 12, "agecat_name": "2014-2015"}
    await store.save_registration(uid, "+998901234567", athlete, extra)
    rows = await store.list_registrations(uid)
    print("saved rows:", len(rows))
    r = rows[0]
    print("fullname:", r["fullname"], "| phone:", r["phone"], "| created:", r["created_at"])
    print("data_json keys:", sorted(r["data_json"].keys())[:12], "...")
    data = exporter.build_excel(rows)
    print("excel bytes:", len(data))
    ws = load_workbook(io_bytes(data)).active
    hdr = [c.value for c in ws[1]]
    row2 = [c.value for c in ws[2]]
    print("cols:", len(hdr), "| row2[0:6]:", row2[0:6], "| ro'yxat vaqti:", row2[-1])


def io_bytes(b):
    import io
    return io.BytesIO(b)


asyncio.run(t())
print("INTEGRATION_OK")